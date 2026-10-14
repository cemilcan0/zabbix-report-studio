# -*- coding: utf-8 -*-
"""
Scheduled report helpers — no Flask or database code in here.

app.py owns the schedules and fetches the events; this module turns a list of
enriched events into numbers (summarize), a workbook (build_excel) and an e-mail
(build_email_html / send_mail), and works out when a schedule runs next and
which period it covers. Text is produced in English or Turkish (LABELS).
"""
import io
import os
import smtplib
import ssl
from collections import Counter
from datetime import datetime, timedelta
from email.message import EmailMessage

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FREQUENCIES = ("daily", "weekly", "monthly")
IN_HOURS = "Mesai İçi"          # value stored in the event cache (see run_report_job)

LABELS = {
    "en": {
        "report": "Alarm Report", "summary": "Summary", "top_hosts": "Top Hosts",
        "top_problems": "Top Problems", "alarms": "Alarms",
        "server": "Server", "period": "Period", "previous_period": "Previous period",
        "host_groups": "Host groups", "all_groups": "All host groups",
        "generated": "Generated", "metric": "Metric", "current": "This period",
        "previous": "Previous", "change": "Change",
        "total": "Total alarms", "resolved": "Resolved", "open": "Still open",
        "in_hours": "Business hours", "after_hours": "After hours",
        "avg_resolve": "Average resolution time", "median_resolve": "Median resolution time",
        "hosts_affected": "Affected hosts",
        "host": "Host", "problem": "Problem", "count": "Alarm count",
        "time": "Time", "severity": "Severity", "status": "Status",
        "duration": "Duration", "shift": "Shift", "resolved_at": "Resolved at",
        "sev": ["Not classified", "Information", "Warning", "Average", "High", "Disaster"],
        "freq": {"daily": "Daily", "weekly": "Weekly", "monthly": "Monthly"},
        "attached": "The full list of alarms is attached as an Excel file, with this summary as a PDF.",
        "footer": "Sent automatically by Zabbix Report Studio.",
        "no_alarms": "No alarms in this period.",
        "test_subject": "Zabbix Report Studio — test e-mail",
        "test_body": "The SMTP settings work: this message was sent by Zabbix Report Studio.",
        "location": "Location", "serial": "Serial Number", "name": "Name",
        "hostname": "Hostname", "ip": "IP Address",
        "enabled": "Enabled", "disabled": "Disabled",
        "no_hosts": "No hosts found for this category.", "all_devices": "All Devices",
        "units": ("d", "h", "m", "s"),
    },
    "tr": {
        "report": "Alarm Raporu", "summary": "Özet", "top_hosts": "En Çok Alarm Veren Hostlar",
        "top_problems": "En Sık Problemler", "alarms": "Alarmlar",
        "server": "Sunucu", "period": "Dönem", "previous_period": "Önceki dönem",
        "host_groups": "Host grupları", "all_groups": "Tüm host grupları",
        "generated": "Oluşturulma", "metric": "Metrik", "current": "Bu dönem",
        "previous": "Önceki", "change": "Değişim",
        "total": "Toplam alarm", "resolved": "Çözülen", "open": "Hâlâ açık",
        "in_hours": "Mesai İçi", "after_hours": "Mesai Dışı",
        "avg_resolve": "Ortalama çözüm süresi", "median_resolve": "Medyan çözüm süresi",
        "hosts_affected": "Etkilenen host",
        "host": "Host", "problem": "Problem", "count": "Alarm sayısı",
        "time": "Zaman", "severity": "Önem", "status": "Durum",
        "duration": "Süre", "shift": "Mesai", "resolved_at": "Çözülme zamanı",
        "sev": ["Sınıflandırılmamış", "Bilgi", "Uyarı", "Orta", "Yüksek", "Felaket"],
        "freq": {"daily": "Günlük", "weekly": "Haftalık", "monthly": "Aylık"},
        "attached": "Alarmların tam listesi Excel dosyası olarak, bu özet de PDF olarak ektedir.",
        "footer": "Zabbix Report Studio tarafından otomatik gönderilmiştir.",
        "no_alarms": "Bu dönemde alarm yok.",
        "test_subject": "Zabbix Report Studio — test e-postası",
        "test_body": "SMTP ayarları çalışıyor: bu ileti Zabbix Report Studio tarafından gönderildi.",
        "location": "Lokasyon", "serial": "Seri Numarası", "name": "Ad",
        "hostname": "Hostname", "ip": "IP Adresi",
        "enabled": "Aktif", "disabled": "Pasif",
        "no_hosts": "Bu kategori için host bulunamadı.", "all_devices": "Tüm Cihazlar",
        "units": ("g", "sa", "dk", "sn"),
    },
}


def labels(lang):
    return LABELS.get(lang) or LABELS["en"]


# ─── Schedule arithmetic ────────────────────────────────
def _month_day(year, month, day):
    """`day` of the month, or its last day when the month is shorter."""
    first_next = datetime(year + month // 12, month % 12 + 1, 1)
    return min(day, (first_next - timedelta(days=1)).day)


def next_run(frequency, hour, weekday=0, monthday=1, after=None):
    """First run time strictly after `after` (local time).
    weekday: 0 = Monday … 6 = Sunday. monthday: 1–31 (clamped to short months)."""
    after = after or datetime.now()
    hour = max(0, min(23, int(hour)))
    base = after.replace(hour=hour, minute=0, second=0, microsecond=0)
    if frequency == "weekly":
        cand = base + timedelta(days=(int(weekday) - base.weekday()) % 7)
        return cand if cand > after else cand + timedelta(days=7)
    if frequency == "monthly":
        year, month = after.year, after.month
        for _ in range(3):
            cand = datetime(year, month, _month_day(year, month, int(monthday)), hour)
            if cand > after:
                return cand
            year, month = year + month // 12, month % 12 + 1
    return base if base > after else base + timedelta(days=1)


def period_for(frequency, run_at=None):
    """Period covered by a run and the period before it, as unix timestamps:
    (time_from, time_till, prev_from, prev_till). Periods end at 23:59:59 of the
    day before the run — daily = yesterday, weekly = the last 7 days,
    monthly = the previous calendar month."""
    run_at = run_at or datetime.now()
    end = run_at.replace(hour=0, minute=0, second=0, microsecond=0)   # exclusive
    if frequency == "monthly":
        end = end.replace(day=1)
        start = (end - timedelta(days=1)).replace(day=1)
        prev_start = (start - timedelta(days=1)).replace(day=1)
    else:
        span = timedelta(days=7 if frequency == "weekly" else 1)
        start = end - span
        prev_start = start - span
    ts = lambda d: int(d.timestamp())
    return ts(start), ts(end) - 1, ts(prev_start), ts(start) - 1


# ─── Numbers ────────────────────────────────────────────
def summarize(events):
    """KPIs of one period from enriched events (the rows run_report_job returns)."""
    durations = sorted(int(e.get("duration") or 0) for e in events
                       if e.get("resolved") and int(e.get("duration") or 0) > 0)
    n = len(durations)
    median = 0
    if n:
        median = durations[n // 2] if n % 2 else (durations[n // 2 - 1] + durations[n // 2]) // 2
    resolved = sum(1 for e in events if e.get("resolved"))
    in_hours = sum(1 for e in events if e.get("mesai") == IN_HOURS)
    hosts = Counter(e.get("host") or "?" for e in events)
    return {
        "total": len(events),
        "resolved": resolved,
        "open": len(events) - resolved,
        "in_hours": in_hours,
        "after_hours": len(events) - in_hours,
        "by_severity": dict(Counter(str(e.get("severity", "0")) for e in events)),
        "avg_resolve": int(sum(durations) / n) if n else 0,
        "median_resolve": median,
        "hosts_affected": len(hosts),
        "top_hosts": hosts.most_common(10),
        "top_problems": Counter(e.get("name") or "?" for e in events).most_common(10),
    }


def fmt_duration(seconds, lang="en"):
    seconds = int(seconds or 0)
    if seconds <= 0:
        return "-"
    d, h, m, s = labels(lang)["units"]
    days, rest = divmod(seconds, 86400)
    hours, rest = divmod(rest, 3600)
    minutes, secs = divmod(rest, 60)
    parts = [f"{v}{u}" for v, u in ((days, d), (hours, h), (minutes, m)) if v]
    if secs and not days:
        parts.append(f"{secs}{s}")
    return " ".join(parts)


def fmt_change(cur, prev):
    """Signed change with a percentage: '+12 (+8.3%)'. No percentage from zero."""
    diff = cur - prev
    if not diff:
        return "0"
    out = f"{diff:+d}"
    return out + (f" ({diff / prev * 100:+.1f}%)" if prev else "")


def comparison_rows(cur, prev, lang="en"):
    """[(label, current, previous, change)] — the table shown in the e-mail and
    on the Summary sheet. Durations are formatted, counts stay numeric."""
    L = labels(lang)
    rows = [(L[k], cur[k], prev[k], fmt_change(cur[k], prev[k]))
            for k in ("total", "resolved", "open", "in_hours", "after_hours", "hosts_affected")]
    for k in ("avg_resolve", "median_resolve"):
        diff = cur[k] - prev[k]
        change = "0" if not diff else ("+" if diff > 0 else "-") + fmt_duration(abs(diff), lang)
        rows.append((L[k], fmt_duration(cur[k], lang), fmt_duration(prev[k], lang), change))
    for sev in ("5", "4", "3", "2", "1", "0"):
        c, p = cur["by_severity"].get(sev, 0), prev["by_severity"].get(sev, 0)
        if c or p:
            rows.append((L["sev"][int(sev)], c, p, fmt_change(c, p)))
    return rows


def _fmt_ts(ts):
    return datetime.fromtimestamp(int(ts)).strftime("%d.%m.%Y %H:%M")


def period_text(time_from, time_till):
    return f"{_fmt_ts(time_from)} – {_fmt_ts(time_till)}"


# ─── Excel ──────────────────────────────────────────────
_HDR_FILL = PatternFill("solid", fgColor="1F3864")
_HDR_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
_ODD_FILL = PatternFill("solid", fgColor="DCE6F2")
_SIDE = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)


def _table(ws, first_row, headers, rows, widths):
    for ci, text in enumerate(headers, 1):
        cell = ws.cell(row=first_row, column=ci, value=text)
        cell.fill, cell.font, cell.border = _HDR_FILL, _HDR_FONT, _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for ri, row in enumerate(rows, first_row + 1):
        for ci, value in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.border = _BORDER
            cell.font = Font(name="Calibri", size=10)
            if (ri - first_row) % 2 == 0:
                cell.fill = _ODD_FILL
    for ci, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = width


def build_excel(meta, cur, prev, events, lang="en"):
    """Workbook bytes: Summary (with the previous-period comparison), Top Hosts,
    Top Problems and the full alarm list.
    meta: server, groups (text), time_from, time_till, prev_from, prev_till."""
    L = labels(lang)
    wb = Workbook()
    ws = wb.active
    ws.title = L["summary"]
    ws["A1"] = f'{L["report"]} — {meta["server"]}'
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    info = [
        (L["period"], period_text(meta["time_from"], meta["time_till"])),
        (L["previous_period"], period_text(meta["prev_from"], meta["prev_till"])),
        (L["host_groups"], meta.get("groups") or L["all_groups"]),
        (L["generated"], datetime.now().strftime("%d.%m.%Y %H:%M")),
    ]
    for ri, (key, value) in enumerate(info, 3):
        ws.cell(row=ri, column=1, value=key).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(row=ri, column=2, value=value).font = Font(name="Calibri", size=10)
    _table(ws, 8, [L["metric"], L["current"], L["previous"], L["change"]],
           comparison_rows(cur, prev, lang), [30, 34, 34, 22])

    for title, key, head in ((L["top_hosts"], "top_hosts", L["host"]),
                             (L["top_problems"], "top_problems", L["problem"])):
        sheet = wb.create_sheet(title=title[:31])
        _table(sheet, 1, [head, L["count"]], cur[key], [70, 16])

    sheet = wb.create_sheet(title=L["alarms"])
    sheet.freeze_panes = "A2"
    rows = []
    for e in events:
        sev = int(e.get("severity") or 0)
        rows.append((
            _fmt_ts(e["clock"]), e.get("host", ""), e.get("name", ""),
            L["sev"][sev] if 0 <= sev <= 5 else str(sev),
            L["resolved"] if e.get("resolved") else L["open"],
            fmt_duration(e.get("duration"), lang) if e.get("resolved") else "-",
            L["in_hours"] if e.get("mesai") == IN_HOURS else L["after_hours"],
        ))
    _table(sheet, 1, [L["time"], L["host"], L["problem"], L["severity"], L["status"],
                      L["duration"], L["shift"]], rows, [18, 28, 70, 18, 14, 14, 14])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── PDF ────────────────────────────────────────────────
# (regular, bold) fonts that cover Turkish; the first pair found is used.
_PDF_FONTS = [
    (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]
_ASCII = str.maketrans("ığüşöçİĞÜŞÖÇâ–—", "igusocIGUSOCa--")


def build_pdf(meta, cur, prev, lang="en"):
    """One-page PDF summary (comparison table, top hosts, top problems) as bytes.
    Returns None when the optional fpdf2 package is not installed. Without a
    Unicode font on the machine, Turkish letters are written without accents."""
    try:
        from fpdf import FPDF
    except ImportError:
        return None
    L = labels(lang)
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=14)
    pdf.add_page()
    fonts = next((pair for pair in _PDF_FONTS if all(os.path.exists(f) for f in pair)), None)
    if fonts:
        pdf.add_font("body", "", fonts[0])
        pdf.add_font("body", "B", fonts[1])
        family, text = "body", str
    else:
        family = "Helvetica"
        text = lambda v: str(v).translate(_ASCII).encode("latin-1", "replace").decode("latin-1")

    def line(value, size=9, style="", gap=5):
        pdf.set_font(family, style, size)
        pdf.multi_cell(0, gap, text(value), new_x="LMARGIN", new_y="NEXT")

    def table(title, headers, rows, widths):
        if title:
            pdf.ln(4)
            line(title, 11, "B", 7)
        pdf.set_font(family, "", 8.5)
        with pdf.table(col_widths=widths, line_height=6, text_align="LEFT") as tbl:
            for row in [headers] + [list(r) for r in rows]:
                cells = tbl.row()
                for value in row:
                    cells.cell(text(value))

    line(f'{L["report"]} — {meta["server"]}', 15, "B", 9)
    line(f'{L["period"]}: {period_text(meta["time_from"], meta["time_till"])}')
    line(f'{L["previous_period"]}: {period_text(meta["prev_from"], meta["prev_till"])}')
    line(f'{L["host_groups"]}: {meta.get("groups") or L["all_groups"]}')
    pdf.ln(2)
    table("", [L["metric"], L["current"], L["previous"], L["change"]],
          comparison_rows(cur, prev, lang), (40, 22, 22, 26))
    if cur["total"]:
        table(L["top_hosts"], [L["host"], L["count"]], cur["top_hosts"], (85, 25))
        table(L["top_problems"], [L["problem"], L["count"]], cur["top_problems"], (85, 25))
    else:
        pdf.ln(4)
        line(L["no_alarms"])
    pdf.ln(5)
    pdf.set_text_color(130, 130, 130)
    line(L["footer"], 7.5)
    return bytes(pdf.output())


# ─── E-mail ─────────────────────────────────────────────
def _esc(value):
    return (str(value).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


def build_email_html(meta, cur, prev, lang="en"):
    L = labels(lang)
    td = 'style="padding:7px 12px;border-bottom:1px solid #e5e5e5;font-size:13px"'
    th = 'style="padding:8px 12px;background:#1f3864;color:#fff;font-size:12px;text-align:left"'

    def table(headers, rows):
        head = "".join(f"<th {th}>{_esc(h)}</th>" for h in headers)
        body = "".join("<tr>" + "".join(f"<td {td}>{_esc(c)}</td>" for c in r) + "</tr>" for r in rows)
        return (f'<table cellspacing="0" cellpadding="0" style="border-collapse:collapse;'
                f'margin:8px 0 22px;min-width:460px"><tr>{head}</tr>{body}</table>')

    parts = [
        '<div style="font-family:Segoe UI,Arial,sans-serif;color:#222;max-width:720px">',
        f'<h2 style="margin:0 0 6px">{_esc(L["report"])} — {_esc(meta["server"])}</h2>',
        f'<p style="margin:0 0 4px;color:#555;font-size:13px"><b>{_esc(L["period"])}:</b> '
        f'{_esc(period_text(meta["time_from"], meta["time_till"]))}</p>',
        f'<p style="margin:0 0 18px;color:#555;font-size:13px"><b>{_esc(L["host_groups"])}:</b> '
        f'{_esc(meta.get("groups") or L["all_groups"])}</p>',
        table([L["metric"], L["current"], L["previous"], L["change"]], comparison_rows(cur, prev, lang)),
    ]
    if cur["total"]:
        parts.append(f'<h3 style="margin:0 0 4px;font-size:15px">{_esc(L["top_hosts"])}</h3>')
        parts.append(table([L["host"], L["count"]], cur["top_hosts"][:5]))
        parts.append(f'<h3 style="margin:0 0 4px;font-size:15px">{_esc(L["top_problems"])}</h3>')
        parts.append(table([L["problem"], L["count"]], cur["top_problems"][:5]))
        parts.append(f'<p style="font-size:13px">{_esc(L["attached"])}</p>')
    else:
        parts.append(f'<p style="font-size:13px">{_esc(L["no_alarms"])}</p>')
    parts.append(f'<p style="color:#999;font-size:11px;margin-top:22px">{_esc(L["footer"])}</p></div>')
    return "".join(parts)


def parse_recipients(text):
    """Split a comma/semicolon/whitespace separated list; returns (valid, invalid)."""
    valid, invalid = [], []
    for item in (text or "").replace(";", ",").replace("\n", ",").replace(" ", ",").split(","):
        item = item.strip()
        if not item:
            continue
        name, at, domain = item.partition("@")
        ok = at and name and "." in domain and not domain.startswith(".") and not domain.endswith(".")
        (valid if ok else invalid).append(item)
    return valid, invalid


_MIME = {'.xlsx': 'vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.pdf': 'pdf'}


def send_mail(smtp, recipients, subject, html, attachments=(), timeout=30):
    """Send an HTML e-mail. smtp: host, port, security ('starttls' | 'ssl' | 'none'),
    username, password, from_addr. attachments: [(filename, bytes)] (.xlsx / .pdf).
    TLS connections verify the server certificate."""
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = smtp["from_addr"]
    msg["To"] = ", ".join(recipients)
    msg.set_content("This message is best viewed in an HTML-capable mail client.")
    msg.add_alternative(html, subtype="html")
    for filename, content in attachments:
        msg.add_attachment(content, maintype="application",
                           subtype=_MIME.get(os.path.splitext(filename)[1], "octet-stream"),
                           filename=filename)
    host, port, security = smtp["host"], int(smtp["port"]), smtp.get("security") or "starttls"
    if security == "ssl":
        server = smtplib.SMTP_SSL(host, port, timeout=timeout, context=ssl.create_default_context())
    else:
        server = smtplib.SMTP(host, port, timeout=timeout)
    try:
        if security == "starttls":
            server.starttls(context=ssl.create_default_context())
        if smtp.get("username"):
            server.login(smtp["username"], smtp.get("password") or "")
        server.send_message(msg)
    finally:
        try:
            server.quit()
        except Exception:
            pass
