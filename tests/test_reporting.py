# -*- coding: utf-8 -*-
import io
from datetime import datetime as D

import pytest
from openpyxl import load_workbook

import reporting as r

ts = lambda *a: int(D(*a).timestamp())
dt = D.fromtimestamp


@pytest.mark.parametrize("args, after, expected", [
    (("daily", 8), D(2026, 10, 13, 7, 59), D(2026, 10, 13, 8)),
    (("daily", 8), D(2026, 10, 13, 8, 0), D(2026, 10, 14, 8)),          # strictly after
    (("weekly", 8, 0), D(2026, 10, 13, 9), D(2026, 10, 19, 8)),         # Tuesday -> next Monday
    (("weekly", 8, 1), D(2026, 10, 13, 7), D(2026, 10, 13, 8)),         # later the same day
    (("weekly", 8, 1), D(2026, 10, 13, 8), D(2026, 10, 20, 8)),
    (("monthly", 6, 0, 31), D(2026, 2, 1), D(2026, 2, 28, 6)),          # short month
    (("monthly", 6, 0, 31), D(2026, 2, 28, 6), D(2026, 3, 31, 6)),
    (("monthly", 6, 0, 1), D(2026, 12, 15), D(2027, 1, 1, 6)),          # year rollover
])
def test_next_run(args, after, expected):
    assert r.next_run(*args, after=after) == expected


def test_period_for():
    f = lambda freq, at: tuple(dt(x) for x in r.period_for(freq, at))
    assert f("daily", D(2026, 10, 13, 8)) == (
        D(2026, 10, 12), D(2026, 10, 12, 23, 59, 59), D(2026, 10, 11), D(2026, 10, 11, 23, 59, 59))
    assert f("weekly", D(2026, 10, 19, 8))[:3] == (D(2026, 10, 12), D(2026, 10, 18, 23, 59, 59), D(2026, 10, 5))
    assert f("monthly", D(2026, 3, 1, 6)) == (
        D(2026, 2, 1), D(2026, 2, 28, 23, 59, 59), D(2026, 1, 1), D(2026, 1, 31, 23, 59, 59))
    assert f("monthly", D(2026, 1, 15, 6))[:3] == (D(2025, 12, 1), D(2025, 12, 31, 23, 59, 59), D(2025, 11, 1))


def test_parse_recipients():
    assert r.parse_recipients("a@b.co; bad, c@d.org\nx@y") == (["a@b.co", "c@d.org"], ["bad", "x@y"])
    assert r.parse_recipients("") == ([], [])


def test_fmt_change():
    assert r.fmt_change(10, 8) == "+2 (+25.0%)"
    assert r.fmt_change(6, 8) == "-2 (-25.0%)"
    assert r.fmt_change(3, 0) == "+3"          # no percentage from zero
    assert r.fmt_change(5, 5) == "0"


EVENTS = [
    {"clock": ts(2026, 10, 5, 10), "host": "a", "name": "CPU", "severity": "4", "resolved": True, "duration": 600, "mesai": r.IN_HOURS},
    {"clock": ts(2026, 10, 5, 22), "host": "a", "name": "CPU", "severity": "2", "resolved": True, "duration": 1200, "mesai": "Mesai Dışı"},
    {"clock": ts(2026, 10, 6, 3), "host": "b", "name": "Disk", "severity": "4", "resolved": False, "duration": 0, "mesai": "Mesai Dışı"},
]
META = {"server": "Test Zabbix", "groups": "", "time_from": ts(2026, 10, 5), "time_till": ts(2026, 10, 11, 23, 59, 59),
        "prev_from": ts(2026, 9, 28), "prev_till": ts(2026, 10, 4, 23, 59, 59)}


def test_summarize():
    s = r.summarize(EVENTS)
    assert (s["total"], s["resolved"], s["open"], s["in_hours"], s["after_hours"]) == (3, 2, 1, 1, 2)
    assert (s["avg_resolve"], s["median_resolve"], s["hosts_affected"]) == (900, 900, 2)
    assert s["by_severity"] == {"4": 2, "2": 1}
    assert s["top_hosts"][0] == ("a", 2) and s["top_problems"][0] == ("CPU", 2)
    empty = r.summarize([])
    assert empty["total"] == 0 and empty["avg_resolve"] == 0 and empty["top_hosts"] == []


def test_comparison_rows_cover_severities_present_in_either_period():
    rows = r.comparison_rows(r.summarize(EVENTS), r.summarize(EVENTS[:1]), "en")
    by_label = {row[0]: row[1:] for row in rows}
    assert by_label["Total alarms"] == (3, 1, "+2 (+200.0%)")
    assert by_label["Average resolution time"] == ("15m", "10m", "+5m")
    assert "Warning" in by_label and "Disaster" not in by_label


@pytest.mark.parametrize("lang, summary, alarms, total_label", [
    ("en", "Summary", "Alarms", "Total alarms"), ("tr", "Özet", "Alarmlar", "Toplam alarm")])
def test_build_excel(lang, summary, alarms, total_label):
    cur, prev = r.summarize(EVENTS), r.summarize([])
    wb = load_workbook(io.BytesIO(r.build_excel(META, cur, prev, EVENTS, lang)))
    assert wb.sheetnames[0] == summary and wb.sheetnames[-1] == alarms
    assert wb[summary]["A9"].value == total_label and wb[summary]["B9"].value == 3
    assert wb[alarms].max_row == 1 + len(EVENTS)


def test_build_pdf():
    pytest.importorskip("fpdf")
    pdf = r.build_pdf(META, r.summarize(EVENTS), r.summarize([]), "tr")
    assert pdf.startswith(b"%PDF") and len(pdf) > 1500


def test_email_html_escapes_names():
    meta = dict(META, server="<b>srv</b>")
    html = r.build_email_html(meta, r.summarize(EVENTS), r.summarize([]), "en")
    assert "&lt;b&gt;srv&lt;/b&gt;" in html and "<b>srv</b>" not in html
