# -*- coding: utf-8 -*-
import email
import io
import sqlite3
import time
from datetime import datetime
from email import policy

import pytest
import requests
from openpyxl import load_workbook

import app as zrs
from conftest import add_server, add_user, run_report, sign_in


# ─── Passwords ───────────────────────────────────────────
def test_passwords_are_stored_hashed_and_not_shown():
    add_user("bob", "BobSecret12")
    with zrs.app.app_context():
        stored = zrs.User.query.filter_by(username="bob").first().password
    assert stored.startswith(("scrypt:", "pbkdf2:")) and "BobSecret12" not in stored
    page = sign_in().get("/users").get_data(as_text=True)
    assert "bob" in page and stored not in page and "BobSecret12" not in page


def test_plaintext_rows_of_an_old_database_are_upgraded():
    con = sqlite3.connect(zrs._USERS_DB)
    con.execute("INSERT INTO users (username, password, role, must_change_password) VALUES ('old', 'OldPlain123', 'user', 0)")
    con.execute("UPDATE users SET password='Admin123!' WHERE username='admin'")
    con.commit()
    con.close()
    with zrs.app.app_context():
        zrs._migrate_schema()
        old = zrs.User.query.filter_by(username="old").first()
        admin = zrs.User.query.filter_by(username="admin").first()
        assert old.password_is_hashed() and old.check_password("OldPlain123") and not old.must_change_password
        assert admin.password_is_hashed() and admin.must_change_password     # the well-known default
    sign_in("old", "OldPlain123")


def test_flagged_user_can_do_nothing_until_the_password_is_changed():
    add_user("temp", "TempPass123", must_change=True)
    client = zrs.app.test_client()
    r = client.post("/login", data={"username": "temp", "password": "TempPass123"})
    assert r.headers["Location"].endswith("/change-password")
    assert client.get("/").headers["Location"].endswith("/change-password")
    api = client.get("/api/servers")
    assert api.status_code == 403 and api.get_json()["mustChangePassword"] is True

    change = lambda cur, new, again=None: client.post("/change-password", data={
        "current_password": cur, "new_password": new, "confirm_password": again or new})
    assert b"current password is wrong" in change("nope", "NewPassword1").data
    assert b"at least 8 characters" in change("TempPass123", "short").data
    assert b"do not match" in change("TempPass123", "NewPassword1", "Different12").data
    assert b"must be different" in change("TempPass123", "TempPass123").data
    assert b"default password cannot be used" in change("TempPass123", zrs.DEFAULT_ADMIN_PASSWORD).data
    assert change("TempPass123", "NewPassword1").status_code == 302
    assert client.get("/api/servers").status_code == 200
    sign_in("temp", "NewPassword1")


def test_admin_typed_passwords_are_temporary():
    admin = sign_in()
    assert b"at least 8 characters" in admin.post("/users/create", data={"username": "eve", "password": "short", "role": "user"}, follow_redirects=True).data
    assert b"Unknown role." in admin.post("/users/create", data={"username": "eve", "password": "TempPass123", "role": "root"}, follow_redirects=True).data
    admin.post("/users/create", data={"username": "eve", "password": "TempPass123", "role": "user"})
    client = zrs.app.test_client()
    assert client.post("/login", data={"username": "eve", "password": "TempPass123"}).headers["Location"].endswith("/change-password")

    user_id = add_user("carl", "CarlPass123")
    carl = sign_in("carl", "CarlPass123")
    assert carl.post(f"/users/reset/{user_id}", data={"password": "Hijacked123"}).status_code == 302     # not an admin
    sign_in("carl", "CarlPass123")
    admin.post(f"/users/reset/{user_id}", data={"password": "ResetPass123"})
    assert zrs.app.test_client().post("/login", data={"username": "carl", "password": "ResetPass123"}).headers["Location"].endswith("/change-password")


def test_sign_in_attempts_are_limited():
    add_user("dora", "DoraPass123")
    client = zrs.app.test_client()
    wrong = lambda: client.post("/login", data={"username": "dora", "password": "wrong"})
    for _ in range(zrs.LOGIN_MAX_PER_USER):
        assert wrong().status_code == 200
    assert wrong().status_code == 429
    blocked = client.post("/login", data={"username": "dora", "password": "DoraPass123"})
    assert blocked.status_code == 429 and b"Too many failed attempts" in blocked.data
    # another account from the same address is not locked out
    add_user("finn", "FinnPass123")
    sign_in("finn", "FinnPass123")
    # ... and the lock ends when the window has passed
    for key in list(zrs._failed_logins):
        zrs._failed_logins[key] = [t - zrs.LOGIN_WINDOW - 1 for t in zrs._failed_logins[key]]
    sign_in("dora", "DoraPass123")


# ─── CSRF ────────────────────────────────────────────────
@pytest.mark.parametrize("headers, allowed", [
    ({}, True),                                              # not a browser page
    ({"Origin": "http://localhost"}, True),
    ({"Origin": "http://evil.example"}, False),
    ({"Referer": "http://evil.example/page"}, False),
    ({"Sec-Fetch-Site": "cross-site"}, False),
    ({"Sec-Fetch-Site": "same-site", "Origin": "http://localhost"}, False),   # another port/subdomain
    ({"Sec-Fetch-Site": "same-origin", "Origin": "http://localhost"}, True),
])
def test_cross_site_posts_are_refused(headers, allowed):
    admin = sign_in()
    r = admin.post("/users/create", data={"username": "mallory", "password": "Mallory123", "role": "admin"}, headers=headers)
    with zrs.app.app_context():
        created = zrs.User.query.filter_by(username="mallory").first() is not None
    assert created is allowed and (r.status_code == 302) is allowed


def test_no_cors_headers():
    r = sign_in().get("/api/servers", headers={"Origin": "http://evil.example"})
    assert "Access-Control-Allow-Origin" not in r.headers


# ─── Roles ───────────────────────────────────────────────
def test_viewer_is_read_only_and_uses_admin_servers(zabbix_old):
    admin = sign_in()
    server_id = add_server(admin, zabbix_old, "Admin server")
    add_user("user1")
    add_server(sign_in("user1", "UserPass99"), zabbix_old, "Private server")
    add_user("vera", role="viewer")
    vera = sign_in("vera", "UserPass99")

    listing = vera.get("/api/servers").get_json()
    assert listing["canManage"] is False and [s["name"] for s in listing["servers"]] == ["Admin server"]
    body = {"name": "x", "url": "http://x", "username": "x", "password": "x"}
    for path in ("/api/servers", f"/api/servers/{server_id}", f"/api/servers/{server_id}/delete",
                 "/alarm/api/cache-clear", "/api/schedules"):
        assert vera.post(path, json=body).status_code == 403, path
    assert vera.get("/api/schedules").status_code == 403 and vera.get("/api/smtp").status_code == 403
    assert vera.get("/schedules").status_code == 302 and vera.get("/users").status_code == 302
    assert vera.post("/api/business-hours", json={"start": 9, "end": 17, "days": [0]}).status_code == 403
    # ... but reading works
    assert vera.post(f"/api/servers/{server_id}/activate").get_json()["success"]
    assert vera.get("/api/dashboard/stats").get_json()["success"]
    assert "/schedules" not in vera.get("/").get_data(as_text=True)


# ─── Zabbix connection ───────────────────────────────────
def test_certificate_errors_are_reported_and_not_retried(monkeypatch):
    calls = []

    def post(url, **kwargs):
        calls.append(kwargs["verify"])
        raise requests.exceptions.SSLError("certificate verify failed: self-signed certificate")

    monkeypatch.setattr(zrs.req_lib, "post", post)
    admin = sign_in()
    server_id = admin.post("/api/servers", json={"name": "TLS", "url": "https://zabbix.invalid", "username": "a", "password": "b"}).get_json()["id"]
    admin.post(f"/api/servers/{server_id}/activate")
    assert admin.get("/api/servers").get_json()["servers"][0]["verifyTls"] is True      # the default

    result = admin.post("/api/servers/check").get_json()["results"][str(server_id)]
    assert result == {"ok": False, "error": zrs.TLS_ERROR}
    started = time.time()
    login = admin.post("/alarm/api/auto-login", json={})
    assert login.status_code == 401 and login.get_json()["error"] == zrs.TLS_ERROR
    assert time.time() - started < 2            # no retry back-off
    assert calls and all(v is True for v in calls)

    calls.clear()
    admin.post(f"/api/servers/{server_id}", json={"name": "TLS", "url": "https://zabbix.invalid", "username": "a", "verifyTls": False})
    admin.post("/api/servers/check")
    assert calls and all(v is False for v in calls)


@pytest.mark.parametrize("fake", ["zabbix_old", "zabbix_new"])
def test_works_with_body_token_and_bearer_token_servers(fake, request):
    """6.0 takes the session token in the body, 7.2 only as a Bearer header."""
    zabbix = request.getfixturevalue(fake)
    zrs._bearer_auth.clear()
    admin = sign_in()
    server_id = add_server(admin, zabbix)
    assert admin.post("/api/servers/check").get_json()["results"][str(server_id)] == {"ok": True}
    assert admin.get("/api/dashboard/stats").get_json()["data"]["monitoredHosts"] == 5      # ZabbixAPI client
    _, rows = run_report(admin)                                                              # alarm report client
    assert len(rows) > 50
    inventory = admin.post("/api/generate").get_json()
    assert inventory["success"] and inventory["summary"]["TOTAL"] == 5
    assert admin.post("/api/tkodu", json={"code": "cpu"}).get_json()["total"] == 5
    assert zrs._bearer_auth[zabbix.url + "/api_jsonrpc.php"][1] is (fake == "zabbix_new")


def test_wrong_zabbix_password_is_reported(zabbix_old):
    admin = sign_in()
    server_id = add_server(admin, zabbix_old, password="wrong")
    result = admin.post("/api/servers/check").get_json()["results"][str(server_id)]
    assert result["ok"] is False and "incorrect" in result["error"]


# ─── Reports ─────────────────────────────────────────────
def test_report_jobs_belong_to_the_user_who_started_them(zabbix_old):
    admin = sign_in()
    add_server(admin, zabbix_old)
    job_id, rows = run_report(admin)
    assert rows
    add_user("user1")
    other = sign_in("user1", "UserPass99")
    assert other.post("/alarm/api/report-result", json={"jobId": job_id}).status_code == 404
    assert other.post("/alarm/api/report-progress", json={"jobId": job_id}).get_json()["status"] == "unknown"
    assert other.post("/alarm/api/report-cancel", json={"jobId": job_id}).get_json()["success"] is False
    # the same for the Zabbix session of the report page
    session_id = admin.post("/alarm/api/auto-login", json={}).get_json()["sessionId"]
    assert other.post("/alarm/api/hostgroups", json={"sessionId": session_id}).status_code == 401
    assert admin.post("/alarm/api/hostgroups", json={"sessionId": session_id}).get_json()["success"]


def test_clearing_the_cache_only_touches_the_active_server(zabbix_old, zabbix_new):
    admin = sign_in()
    first = add_server(admin, zabbix_old, "First")
    run_report(admin)
    add_server(admin, zabbix_new, "Second")
    run_report(admin)
    stats = lambda: admin.post("/alarm/api/cache-stats", json={}).get_json()["data"]["events"]
    assert stats() > 0
    assert admin.post("/alarm/api/cache-clear", json={}).get_json()["success"]      # "Second" is active
    assert stats() == 0
    admin.post(f"/api/servers/{first}/activate")
    assert stats() > 0


def test_business_hours_setting_changes_the_shift_of_cached_events(zabbix_old):
    admin = sign_in()
    add_server(admin, zabbix_old)
    in_hours = lambda rows: sum(1 for e in rows if e["mesai"] == "Mesai İçi")
    _, before = run_report(admin)
    assert 0 < in_hours(before) < len(before)                       # default: weekdays 08–18

    assert admin.get("/api/business-hours").get_json()["hours"] == zrs.DEFAULT_BUSINESS_HOURS
    for bad in ({"start": 18, "end": 8, "days": [0]}, {"start": 8, "end": 18, "days": []}, {"start": 8, "end": 18, "days": [9]}):
        assert admin.post("/api/business-hours", json=bad).status_code == 400
    assert admin.post("/api/business-hours", json={"start": 0, "end": 24, "days": [0, 1, 2, 3, 4, 5, 6]}).get_json()["success"]
    _, after = run_report(admin)                                    # served from the cache
    assert len(after) == len(before) and in_hours(after) == len(after)
    browse = admin.post("/alarm/api/cache-browse", json={"timeFrom": 0}).get_json()
    assert browse["total"] > 0 and all(e["mesai"] == "Mesai İçi" for e in browse["data"])


def test_mesai_for_uses_days_and_hours():
    at = lambda *a: int(datetime(*a).timestamp())
    assert zrs.mesai_for(at(2026, 10, 12, 8, 0)) == "Mesai İçi"         # Monday 08:00
    assert zrs.mesai_for(at(2026, 10, 12, 18, 0)) == "Mesai Dışı"       # end is exclusive
    assert zrs.mesai_for(at(2026, 10, 10, 10, 0)) == "Mesai Dışı"       # Saturday
    zrs._business_hours.update(start=9, end=13, days=[5])
    assert zrs.mesai_for(at(2026, 10, 10, 10, 0)) == "Mesai İçi"
    assert zrs.mesai_for(at(2026, 10, 12, 10, 0)) == "Mesai Dışı"


def test_on_call_lookup_handles_the_night_after_midnight():
    cfg = {"baslangic": datetime(2026, 10, 12).date(), "dongu_len": 2, "day_start": 8, "night_start": 20,
           "dongu": [{"gunduz": "Ali", "gece": "Veli"}, {"gunduz": "Can", "gece": "Ece"}]}
    at = lambda *a: int(datetime(*a).timestamp())
    assert zrs.nobetci_for(at(2026, 10, 12, 9), cfg) == "Ali"
    assert zrs.nobetci_for(at(2026, 10, 12, 23), cfg) == "Veli"
    assert zrs.nobetci_for(at(2026, 10, 13, 3), cfg) == "Veli"          # still Monday's night shift
    assert zrs.nobetci_for(at(2026, 10, 13, 9), cfg) == "Can"
    assert zrs.nobetci_for(at(2026, 10, 14, 9), cfg) == "Ali"           # the cycle repeats
    assert zrs.nobetci_for(at(2026, 10, 12, 9), None) == ""


@pytest.mark.parametrize("lang, header, sheet, status", [
    ("en", "Location", "All Devices", "Enabled"), ("tr", "Lokasyon", "Tüm Cihazlar", "Aktif"), (None, "Location", "All Devices", "Enabled")])
def test_inventory_excel_follows_the_interface_language(zabbix_old, lang, header, sheet, status):
    admin = sign_in()
    add_server(admin, zabbix_old)
    response = admin.post("/api/download", json={"lang": lang}) if lang else admin.post("/api/download")
    ws = load_workbook(io.BytesIO(response.data)).active
    assert (ws.title, ws["A1"].value, ws["F2"].value) == (sheet, header, status)


# ─── Scheduled reports ───────────────────────────────────
def schedule_body(server_id, **changes):
    return dict({"name": "Weekly", "serverId": server_id, "frequency": "weekly", "weekday": 0, "monthday": 1,
                 "hour": 8, "groups": [{"id": "11", "name": "Linux Servers"}], "severities": ["2", "3", "4", "5"],
                 "recipients": "noc@example.com; boss@example.com", "lang": "tr", "enabled": True}, **changes)


def test_schedule_validation_and_ownership(zabbix_old):
    admin = sign_in()
    server_id = add_server(admin, zabbix_old)
    add_user("user1")
    other = sign_in("user1", "UserPass99")
    post = lambda client, **changes: client.post("/api/schedules", json=schedule_body(server_id, **changes))
    assert "Invalid e-mail address" in post(admin, recipients="nope").get_json()["error"]
    assert post(admin, severities=[]).status_code == 400
    assert post(admin, frequency="hourly").status_code == 400
    assert post(other).get_json()["error"] == "Server not found."          # not their server
    schedule_id = post(admin).get_json()["id"]

    row = admin.get("/api/schedules").get_json()["schedules"][0]
    run_at = time.localtime(row["nextRun"])
    assert row["nextRun"] > time.time() and (run_at.tm_wday, run_at.tm_hour) == (0, 8)
    assert row["recipients"] == "noc@example.com, boss@example.com"
    assert other.get("/api/schedules").get_json()["schedules"] == []
    for action in ("run", "delete"):
        assert other.post(f"/api/schedules/{schedule_id}/{action}").status_code == 404
    assert admin.post(f"/api/schedules/{schedule_id}/run").get_json()["error"] == "SMTP is not configured."
    assert other.post("/api/smtp", json={}).status_code == 403


def test_scheduled_report_is_mailed_with_excel_and_pdf(zabbix_new, smtp):
    admin = sign_in()
    server_id = add_server(admin, zabbix_new, "Prod Zabbix")
    assert admin.post("/api/smtp", json={"host": "", "port": 25, "fromAddr": "zrs@example.com"}).status_code == 400
    assert admin.post("/api/smtp", json={"host": "127.0.0.1", "port": smtp.port, "security": "none",
                                         "fromAddr": "zrs@example.com"}).get_json()["success"]
    assert admin.post("/api/smtp/test", json={"to": "me@example.com"}).get_json()["success"]
    assert len(smtp.messages) == 1

    schedule_id = admin.post("/api/schedules", json=schedule_body(server_id)).get_json()["id"]
    zrs.run_schedule(schedule_id)                       # what the scheduler thread calls
    row = admin.get("/api/schedules").get_json()["schedules"][0]
    assert row["lastStatus"].startswith("Sent to 2 recipients"), row["lastStatus"]
    assert len(smtp.messages) == 2

    message = email.message_from_bytes(smtp.messages[-1], policy=policy.default)
    html = message.get_body(preferencelist=("html",)).get_content()
    files = {a.get_filename().rsplit(".", 1)[1]: a.get_payload(decode=True) for a in message.iter_attachments()}
    assert message["Subject"].startswith("Alarm Raporu — Prod Zabbix — Haftalık")
    assert message["To"] == "noc@example.com, boss@example.com"
    assert "Toplam alarm" in html and "Linux Servers" in html
    assert "xlsx" in files

    wb = load_workbook(io.BytesIO(files["xlsx"]))
    alarms = wb["Alarmlar"]
    total = wb["Özet"]["B9"].value
    assert total == alarms.max_row - 1 and total > 0
    assert row["lastStatus"].endswith(f"— {total} alarms")
    assert {alarms.cell(row=i, column=2).value for i in range(2, alarms.max_row + 1)} <= {"app-01", "app-02", "db-01"}
    assert "Bilgi" not in {alarms.cell(row=i, column=4).value for i in range(2, alarms.max_row + 1)}    # severity 1 not selected
    try:
        import fpdf  # noqa: F401
        assert files["pdf"].startswith(b"%PDF")
    except ImportError:
        assert "pdf" not in files


def test_failed_schedule_run_records_the_error(zabbix_old, smtp):
    admin = sign_in()
    server_id = add_server(admin, zabbix_old, password="wrong")
    admin.post("/api/smtp", json={"host": "127.0.0.1", "port": smtp.port, "security": "none", "fromAddr": "zrs@example.com"})
    schedule_id = admin.post("/api/schedules", json=schedule_body(server_id)).get_json()["id"]
    zrs.run_schedule(schedule_id)
    status = admin.get("/api/schedules").get_json()["schedules"][0]["lastStatus"]
    assert status.startswith("Error:") and "incorrect" in status and len(smtp.messages) == 0


# ─── Pages ───────────────────────────────────────────────
@pytest.mark.parametrize("path", ["/", "/servers", "/alarm/report", "/ack", "/nobet", "/tkodu", "/envanter",
                                  "/alarm/cache", "/users", "/schedules", "/change-password"])
def test_pages_render(path):
    response = sign_in().get(path)
    assert response.status_code == 200 and b"</html>" in response.data


def test_pages_need_a_sign_in():
    client = zrs.app.test_client()
    assert client.get("/").status_code == 302
    assert client.post("/api/servers", json={}).status_code in (302, 401)
